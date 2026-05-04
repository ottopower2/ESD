# -*- coding: utf-8 -*-
# =============================================================================
# Copyright (c) 2026  CSA GROUP. Alle Rechte vorbehalten / All rights reserved.
# Vertraulich und proprietär – Weitergabe, Kopie und Nutzung durch Dritte
# ohne ausdrückliche schriftliche Genehmigung des Urhebers untersagt.
# Confidential and proprietary – redistribution and use by third parties
# without explicit written permission of the author is strictly prohibited.
# Siehe / See: LICENSE
# =============================================================================
"""
ESD-Bericht-Generator  |  ISO 10605 mit 330 pF / 330 Ω.
Erzeugt eine Word-.docx-Datei mit Messparametern, Rohwerten und PASS/FAIL-Auswertung.
"""
import os
from datetime import datetime

from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


# --------------------------------------------------
CRITERIA = {
    1: {"t65": (2,    0.30), "t130": (1,   0.30), "peak": (3.75,  0.10)},
    2: {"t65": (4,    0.30), "t130": (2,   0.30), "peak": (7.5,   0.10)},
    4: {"t65": (8,    0.30), "t130": (4,   0.30), "peak": (15.0,  0.10)},
    6: {"t65": (12,   0.30), "t130": (6,   0.30), "peak": (22.5,  0.10)},
    8: {"t65": (16,   0.30), "t130": (8,   0.30), "peak": (30.0,  0.10)},
}

COLOR_PASS   = "C6EFCE"   # hellgrün
COLOR_FAIL   = "FFC7CE"   # hellrot
COLOR_HEADER = "D9E1F2"   # hellblau


# --------------------------------------------------
def _check(value, target, tol):
    lo, hi = target * (1 - tol), target * (1 + tol)
    return lo, hi, lo <= value <= hi


def _set_cell_bg(cell, hex_color: str):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement("w:shd")
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  hex_color)
    tcPr.append(shd)


def _set_cell_text(cell, text: str, bold=False, color: RGBColor = None,
                   align=WD_ALIGN_PARAGRAPH.LEFT, font_size=10):
    cell.text = ""
    para = cell.paragraphs[0]
    para.alignment = align
    run = para.add_run(text)
    run.font.size = Pt(font_size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color


def _add_header_row(table, labels: list):
    row = table.rows[0]
    for i, label in enumerate(labels):
        _set_cell_text(row.cells[i], label, bold=True,
                       align=WD_ALIGN_PARAGRAPH.CENTER)
        _set_cell_bg(row.cells[i], COLOR_HEADER)


def _add_row(table, values: list, bg_color: str = None):
    row = table.add_row()
    for i, val in enumerate(values):
        align = WD_ALIGN_PARAGRAPH.RIGHT if i > 0 else WD_ALIGN_PARAGRAPH.LEFT
        _set_cell_text(row.cells[i], str(val), align=align)
        if bg_color:
            _set_cell_bg(row.cells[i], bg_color)
    return row


def _add_pf_row(table, label, value, lo, hi, passed):
    color   = COLOR_PASS if passed else COLOR_FAIL
    verdict = "PASS"     if passed else "FAIL"
    row = table.add_row()
    _set_cell_text(row.cells[0], label)
    _set_cell_text(row.cells[1], f"{value:.4f} A",          align=WD_ALIGN_PARAGRAPH.RIGHT)
    _set_cell_text(row.cells[2], f"[{lo:.4f}, {hi:.4f}] A", align=WD_ALIGN_PARAGRAPH.CENTER)
    _set_cell_text(row.cells[3], verdict, bold=True,
                   color=RGBColor(0x27, 0xAE, 0x60) if passed else RGBColor(0xC0, 0x39, 0x2B),
                   align=WD_ALIGN_PARAGRAPH.CENTER)
    for cell in row.cells:
        _set_cell_bg(cell, color)


# --------------------------------------------------
def generate_report(data: dict, save_dir: str) -> str:
    """
    Erstellt einen Word-Prüfbericht (.docx) und gibt den Dateipfad zurück.

    Erwartete Felder in data:
        timestamp       str   z.B. "20260430_143602"
        entlade_kv      int   z.B. 4 oder -4
        trigger_level   int/float
        trigger_offset  int/float
        scale           int/float
        attenuation     int/float
        v_65ns          float  (Rohspannung, kann negativ sein)
        v_130ns         float  (Rohspannung, kann negativ sein)
        vpeak           float
        vpp             float
        A_65ns          float
        A_130ns         float
        A_peak          float
        screenshot_path str
    """
    kv_abs = abs(data["entlade_kv"])
    c      = CRITERIA[kv_abs]

    A_65ns  = data["A_65ns"]
    A_130ns = data["A_130ns"]
    A_peak  = data["A_peak"]

    lo65,  hi65,  pass65  = _check(A_65ns,  *c["t65"])
    lo130, hi130, pass130 = _check(A_130ns, *c["t130"])
    lopk,  hipk,  passpk  = _check(A_peak,  *c["peak"])

    overall = "PASS" if (pass65 and pass130 and passpk) else "FAIL"
    ts      = data["timestamp"]
    dt_str  = datetime.strptime(ts, "%Y%m%d_%H%M%S").strftime("%d.%m.%Y  %H:%M:%S")

    doc = Document()

    # -- Seitenränder --
    for section in doc.sections:
        section.top_margin    = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin   = Cm(2.5)
        section.right_margin  = Cm(2.5)

    # -- Titel --
    title = doc.add_heading("ESD-Prüfbericht — ISO 10605 mit 330 pF / 330 Ω", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()

    # -- Messparameter --
    doc.add_heading("Messparameter", level=2)
    t = doc.add_table(rows=1, cols=2)
    t.style = "Table Grid"
    _add_header_row(t, ["Parameter", "Wert"])
    for label, val in [
        ("Datum / Uhrzeit",  dt_str),
        ("Entladespannung",  f"{data['entlade_kv']} kV"),
        ("Triggerlevel",     f"{data['trigger_level']} V"),
        ("Trigger-Offset",   f"{data['trigger_offset']} V"),
        ("Skala",            f"{data['scale']} V/div"),
        ("Attenuation",      str(data["attenuation"])),
    ]:
        _add_row(t, [label, val])

    doc.add_paragraph()

    # -- Rohwerte --
    doc.add_heading("Rohwerte (Spannung)", level=2)
    t2 = doc.add_table(rows=1, cols=2)
    t2.style = "Table Grid"
    _add_header_row(t2, ["Größe", "Wert"])
    for label, val in [
        ("V @ 65 ns",  f"{data['v_65ns']:.4f} V"),
        ("V @ 130 ns", f"{data['v_130ns']:.4f} V"),
        ("V_peak",     f"{data['vpeak']:.4f} V"),
        ("V_pp",       f"{data['vpp']:.4f} V"),
    ]:
        _add_row(t2, [label, val])

    doc.add_paragraph()

    # -- Berechnete Ströme --
    doc.add_heading("Berechnete Ströme", level=2)
    t3 = doc.add_table(rows=1, cols=3)
    t3.style = "Table Grid"
    _add_header_row(t3, ["Größe", "Formel", "Wert"])
    ATT = data["attenuation"]
    for label, formula, val in [
        ("I @ 65 ns",  f"|V_65ns| × 10 / (ATT × 2)  =  |V_65ns| × 10 / ({ATT}×2)",  f"{A_65ns:.4f} A"),
        ("I @ 130 ns", f"|V_130ns| × 10 / (ATT × 2)       =  |V_130ns| × 10 / ({ATT}×2)",      f"{A_130ns:.4f} A"),
        ("I_Peak",     f"|V_pp| × 10 / (ATT × 2)     =  |V_pp| × 10 / ({ATT}×2)",     f"{A_peak:.4f} A"),
    ]:
        _add_row(t3, [label, formula, val])

    doc.add_paragraph()

    # -- PASS / FAIL --
    doc.add_heading("PASS / FAIL Auswertung", level=2)
    t4 = doc.add_table(rows=1, cols=4)
    t4.style = "Table Grid"
    _add_header_row(t4, ["Kriterium", "Gemessen", "Grenzwerte", "Ergebnis"])
    _add_pf_row(t4,
                f"I @ 65 ns  — Soll {c['t65'][0]} A ±{int(c['t65'][1]*100)} %",
                A_65ns, lo65, hi65, pass65)
    _add_pf_row(t4,
                f"I @ 130 ns — Soll {c['t130'][0]} A ±{int(c['t130'][1]*100)} %",
                A_130ns, lo130, hi130, pass130)
    _add_pf_row(t4,
                f"I_Peak     — Soll {c['peak'][0]} A ±{int(c['peak'][1]*100)} %",
                A_peak, lopk, hipk, passpk)

    doc.add_paragraph()

    # -- Gesamtergebnis --
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(f"Gesamtergebnis:  {overall}")
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = (
        RGBColor(0x27, 0xAE, 0x60) if overall == "PASS" else RGBColor(0xC0, 0x39, 0x2B)
    )

    doc.add_paragraph()

    # -- Aktueller Screenshot --
    doc.add_heading("Oszilloskop-Screenshot (aktuelle Messung)", level=2)
    screenshot = data["screenshot_path"]
    if os.path.isfile(screenshot):
        doc.add_picture(screenshot, width=Inches(6))
    else:
        doc.add_paragraph(f"[Screenshot nicht gefunden: {os.path.basename(screenshot)}]")

    doc.add_paragraph()

    # -- Alle weiteren Screenshots im Verzeichnis --
    all_pngs = sorted(
        [f for f in os.listdir(save_dir)
         if f.startswith("ESD_") and f.endswith(".png")]
    )
    other_pngs = [f for f in all_pngs if os.path.join(save_dir, f) != screenshot]

    if other_pngs:
        doc.add_heading("Alle weiteren Messungen im Verzeichnis", level=2)
        for png in other_pngs:
            png_path = os.path.join(save_dir, png)
            ts_png   = png[4:-4]          # "ESD_20260430_143602.png" → "20260430_143602"
            json_path = os.path.join(save_dir, f"ESD_{ts_png}.json")

            # Dateiname als Überschrift
            p = doc.add_paragraph()
            run = p.add_run(png)
            run.font.bold = True
            run.font.size = Pt(10)

            # JSON vorhanden → Kriterien auswerten
            if os.path.isfile(json_path):
                import json as _json
                with open(json_path, encoding="utf-8") as jf:
                    d = _json.load(jf)

                kv_abs = abs(d["entlade_kv"])
                c_png  = CRITERIA[kv_abs]
                lo65,  hi65,  p65  = _check(d["A_65ns"],  *c_png["t65"])
                lo130, hi130, p130 = _check(d["A_130ns"], *c_png["t130"])
                lopk,  hipk,  ppk  = _check(d["A_peak"],  *c_png["peak"])
                ov = "PASS" if (p65 and p130 and ppk) else "FAIL"

                mini = doc.add_table(rows=1, cols=4)
                mini.style = "Table Grid"
                _add_header_row(mini, ["Kriterium", "Gemessen", "Grenzwerte", "Ergebnis"])
                _add_pf_row(mini,
                    f"I@65ns  Soll {c_png['t65'][0]} A ±{int(c_png['t65'][1]*100)}%",
                    d["A_65ns"], lo65, hi65, p65)
                _add_pf_row(mini,
                    f"I@130ns Soll {c_png['t130'][0]} A ±{int(c_png['t130'][1]*100)}%",
                    d["A_130ns"], lo130, hi130, p130)
                _add_pf_row(mini,
                    f"I_Peak  Soll {c_png['peak'][0]} A ±{int(c_png['peak'][1]*100)}%",
                    d["A_peak"], lopk, hipk, ppk)

                p2 = doc.add_paragraph()
                run2 = p2.add_run(f"Entladespannung: {d['entlade_kv']} kV   |   Gesamtergebnis: {ov}")
                run2.font.bold = True
                run2.font.size = Pt(10)
                run2.font.color.rgb = (
                    RGBColor(0x27, 0xAE, 0x60) if ov == "PASS" else RGBColor(0xC0, 0x39, 0x2B)
                )
            else:
                doc.add_paragraph("[Keine Messdaten gefunden — nur Screenshot]")

            doc.add_picture(png_path, width=Inches(6))
            doc.add_paragraph()

    # -- Footer-Info --
    p = doc.add_paragraph(
        f"Erstellt: {dt_str}   |   Norm: ISO 10605  |  330 pF / 330 Ω   |   "
        f"Entwickelt von: Amouri   |   © 2026 Alle Rechte vorbehalten"
    )
    p.runs[0].font.size  = Pt(8)
    p.runs[0].font.color.rgb = RGBColor(0x7F, 0x8C, 0x8D)

    # -- Speichern --
    report_name = f"ESD_Report_{ts}.docx"
    report_path = os.path.join(save_dir, report_name)
    doc.save(report_path)
    return report_path


# --------------------------------------------------
if __name__ == "__main__":
    import json as _json
    from datetime import datetime as _dt

    save_dir  = os.path.dirname(os.path.abspath(__file__))
    timestamp = _dt.now().strftime("%Y%m%d_%H%M%S")

    # Alle JSON-Messdateien einlesen, chronologisch sortiert
    json_files = sorted(
        [f for f in os.listdir(save_dir) if f.startswith("ESD_") and f.endswith(".json")]
    )

    if not json_files:
        print("Keine Messdaten (ESD_*.json) gefunden.")
        raise SystemExit

    print(f"{len(json_files)} Messung(en) gefunden:")
    for f in json_files:
        print(f"  {f}")

    all_data = []
    for f in json_files:
        with open(os.path.join(save_dir, f), encoding="utf-8") as jf:
            all_data.append(_json.load(jf))

    # --- Word-Dokument aufbauen ---
    doc = Document()
    for section in doc.sections:
        section.top_margin    = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin   = Cm(2.5)
        section.right_margin  = Cm(2.5)

    title = doc.add_heading("ESD-Gesamtprüfbericht — ISO 10605  |  330 pF / 330 Ω", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph(
        f"Erstellt: {_dt.now().strftime('%d.%m.%Y  %H:%M:%S')}   |   "
        f"Anzahl Messungen: {len(all_data)}   |   Entwickelt von: H.Amouri   |   © 2026"
    )
    p.runs[0].font.size = Pt(9)
    p.runs[0].font.color.rgb = RGBColor(0x7F, 0x8C, 0x8D)
    doc.add_paragraph()

    # Übersichtstabelle aller Messungen
    doc.add_heading("Übersicht aller Messungen", level=2)
    t_ov = doc.add_table(rows=1, cols=6)
    t_ov.style = "Table Grid"
    _add_header_row(t_ov, ["Nr.", "Datum / Uhrzeit", "kV", "I@65ns (A)", "I@130ns (A)", "Gesamtergebnis"])

    for i, d in enumerate(all_data, 1):
        kv_abs = abs(d["entlade_kv"])
        c      = CRITERIA[kv_abs]
        _, _, p65  = _check(d["A_65ns"],  *c["t65"])
        _, _, p130 = _check(d["A_130ns"], *c["t130"])
        _, _, ppk  = _check(d["A_peak"],  *c["peak"])
        ov = "PASS" if (p65 and p130 and ppk) else "FAIL"
        row = t_ov.add_row()
        vals = [
            str(i),
            _dt.strptime(d["timestamp"], "%Y%m%d_%H%M%S").strftime("%d.%m.%Y %H:%M:%S"),
            f"{d['entlade_kv']} kV",
            f"{d['A_65ns']:.3f}",
            f"{d['A_130ns']:.3f}",
            ov,
        ]
        for j, val in enumerate(vals):
            _set_cell_text(row.cells[j], val, align=WD_ALIGN_PARAGRAPH.CENTER)
        _set_cell_bg(row.cells[5], COLOR_PASS if ov == "PASS" else COLOR_FAIL)
        row.cells[5].paragraphs[0].runs[0].font.bold = True

    doc.add_paragraph()

    # Detailsektion pro Messung
    for i, d in enumerate(all_data, 1):
        kv_abs = abs(d["entlade_kv"])
        c      = CRITERIA[kv_abs]
        A_65ns  = d["A_65ns"]
        A_130ns = d["A_130ns"]
        A_peak  = d["A_peak"]
        lo65,  hi65,  pass65  = _check(A_65ns,  *c["t65"])
        lo130, hi130, pass130 = _check(A_130ns, *c["t130"])
        lopk,  hipk,  passpk  = _check(A_peak,  *c["peak"])
        ov     = "PASS" if (pass65 and pass130 and passpk) else "FAIL"
        dt_str = _dt.strptime(d["timestamp"], "%Y%m%d_%H%M%S").strftime("%d.%m.%Y  %H:%M:%S")

        doc.add_heading(f"Messung {i}  —  {d['entlade_kv']} kV  |  {dt_str}", level=2)

        # Messparameter
        t_mp = doc.add_table(rows=1, cols=2)
        t_mp.style = "Table Grid"
        _add_header_row(t_mp, ["Parameter", "Wert"])
        for label, val in [
            ("Entladespannung",  f"{d['entlade_kv']} kV"),
            ("Triggerlevel",     f"{d['trigger_level']} V"),
            ("Trigger-Offset",   f"{d['trigger_offset']} V"),
            ("Skala",            f"{d['scale']} V/div"),
            ("Attenuation",      str(d["attenuation"])),
            ("V @ 65 ns",        f"{d['v_65ns']:.4f} V"),
            ("V @ 130 ns",       f"{d['v_130ns']:.4f} V"),
            ("V_pp",             f"{d['vpp']:.4f} V"),
        ]:
            _add_row(t_mp, [label, val])

        doc.add_paragraph()

        # PASS/FAIL
        t_pf = doc.add_table(rows=1, cols=4)
        t_pf.style = "Table Grid"
        _add_header_row(t_pf, ["Kriterium", "Gemessen", "Grenzwerte", "Ergebnis"])
        _add_pf_row(t_pf, f"I@65ns   Soll {c['t65'][0]} A ±{int(c['t65'][1]*100)}%",
                    A_65ns,  lo65,  hi65,  pass65)
        _add_pf_row(t_pf, f"I@130ns  Soll {c['t130'][0]} A ±{int(c['t130'][1]*100)}%",
                    A_130ns, lo130, hi130, pass130)
        _add_pf_row(t_pf, f"I_Peak   Soll {c['peak'][0]} A ±{int(c['peak'][1]*100)}%",
                    A_peak,  lopk,  hipk,  passpk)

        doc.add_paragraph()

        # Gesamtergebnis
        p2 = doc.add_paragraph()
        run = p2.add_run(f"Gesamtergebnis: {ov}")
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = (
            RGBColor(0x27, 0xAE, 0x60) if ov == "PASS" else RGBColor(0xC0, 0x39, 0x2B)
        )

        # Screenshot
        png = d["screenshot_path"]
        if os.path.isfile(png):
            doc.add_picture(png, width=Inches(6))
        else:
            doc.add_paragraph(f"[Screenshot nicht gefunden: {os.path.basename(png)}]")

        doc.add_paragraph()

    report_name = f"ESD_Gesamtbericht_{timestamp}.docx"
    report_path = os.path.join(save_dir, report_name)
    doc.save(report_path)
    print(f"Gesamtbericht gespeichert: {report_path}")
